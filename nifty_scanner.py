#!/usr/bin/env python3
"""
NIFTY STRUCTURAL SCANNER  -  Volume Profile + Multi-Anchor AVWAP + Signal Ranking
==================================================================================
One run does everything, end to end:

  1. Pulls LIVE constituents of every Nifty index from niftyindices.com
     (falls back to the last good copy in ./cache, then to built-in Nifty 50).
  2. Downloads ~2y of daily OHLCV for the whole universe via yfinance (batched, retried).
  3. Computes per stock: 50-bin Volume Profile (POC / VAH / VAL, 70% value area),
     Anchored VWAPs (Catalyst / Earnings / Swing-Low), squeeze status, structural signal.
  4. Scores every stock 0-100 and labels it STRONG BUY / BUY / WATCH / NEUTRAL / AVOID.
  5. Writes a formatted Excel dashboard with an Index dropdown, Sort dropdown and Signal filter.
     The ranking is done by live Excel formulas, so changing a dropdown re-sorts instantly.

Usage
-----
  python nifty_scanner.py                          # live run (do this AFTER 15:30 IST)
  python nifty_scanner.py --seed-xlsx old_file.xlsx --out seed.xlsx   # build from an existing snapshot
  python nifty_scanner.py --demo --out /tmp/demo.xlsx                 # offline self-test, synthetic prices

Inputs (lookback, catalyst / earnings / swing-low dates) are read from the Scanner sheet of the
existing output file, so you edit them in Excel and the next EOD run picks them up.

Not investment advice. Signals are rule-based technical screens on delayed public data.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import logging
import os
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd

LOG = logging.getLogger("scanner")
HERE = Path(__file__).resolve().parent
CACHE = HERE / "cache"

# ----------------------------------------------------------------------------- config
ALL = "All Indices (Universe)"

# index name -> constituent CSV published by NSE Indices. File names follow NSE's convention;
# any that fail are reported in the Run Log sheet and the cached copy is used instead.
INDEX_FILES = {
    "Nifty 50": "ind_nifty50list.csv",
    "Nifty Next 50": "ind_niftynext50list.csv",
    "Nifty 100": "ind_nifty100list.csv",
    "Nifty 200": "ind_nifty200list.csv",
    "Nifty 500": "ind_nifty500list.csv",
    "Nifty Midcap 100": "ind_niftymidcap100list.csv",
    "Nifty Midcap 150": "ind_niftymidcap150list.csv",
    "Nifty Smallcap 100": "ind_niftysmallcap100list.csv",
    "Nifty Smallcap 250": "ind_niftysmallcap250list.csv",
    "Nifty Bank": "ind_niftybanklist.csv",
    "Nifty Private Bank": "ind_nifty_privatebanklist.csv",
    "Nifty PSU Bank": "ind_niftypsubanklist.csv",
    "Nifty Financial Services": "ind_niftyfinancelist.csv",
    "Nifty IT": "ind_niftyitlist.csv",
    "Nifty Pharma": "ind_niftypharmalist.csv",
    "Nifty Healthcare": "ind_niftyhealthcarelist.csv",
    "Nifty Auto": "ind_niftyautolist.csv",
    "Nifty FMCG": "ind_niftyfmcglist.csv",
    "Nifty Metal": "ind_niftymetallist.csv",
    "Nifty Energy": "ind_niftyenergylist.csv",
    "Nifty Oil & Gas": "ind_niftyoilgaslist.csv",
    "Nifty Realty": "ind_niftyrealtylist.csv",
    "Nifty Media": "ind_niftymedialist.csv",
    "Nifty Consumer Durables": "ind_niftyconsumerdurableslist.csv",
    "Nifty Infrastructure": "ind_niftyinfralist.csv",
    "Nifty Commodities": "ind_niftycommoditieslist.csv",
}
CSV_BASES = [
    "https://niftyindices.com/IndexConstituent/",
    "https://www.niftyindices.com/IndexConstituent/",
    "https://nsearchives.nseindia.com/content/indices/",
]
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/120 Safari/537.36"}

NIFTY50_FALLBACK = (
    "RELIANCE TCS HDFCBANK ICICIBANK INFY BHARTIARTL ITC LT SBIN KOTAKBANK AXISBANK HINDUNILVR "
    "BAJFINANCE MARUTI M&M SUNPHARMA ULTRACEMCO NTPC TITAN POWERGRID TATAMOTORS TATASTEEL ASIANPAINT "
    "ADANIENT ADANIPORTS WIPRO HCLTECH TECHM ONGC COALINDIA NESTLEIND BAJAJFINSV JSWSTEEL GRASIM "
    "HINDALCO DRREDDY CIPLA APOLLOHOSP BRITANNIA EICHERMOT HEROMOTOCO BAJAJ-AUTO DIVISLAB TATACONSUM "
    "SHRIRAMFIN SBILIFE HDFCLIFE INDUSINDBK TRENT BEL"
).split()

DEFAULTS = dict(
    index="Nifty 50", sort="Signal tier, then Score", signal_filter="All Signals",
    lookback=60, catalyst=dt.date(2026, 1, 2), earnings=dt.date(2026, 7, 15), swing_low=dt.date(2026, 8, 5),
)
BINS, VALUE_AREA, SQUEEZE_BAND, VOL_AVG_DAYS = 50, 0.70, 0.02, 20
SORT_OPTIONS = ["Signal tier, then Score", "Volume x Avg", "1D % Change", "5D % Return", "Tightest Squeeze"]
SIGNALS = ["STRONG BUY", "BUY", "WATCH", "NEUTRAL", "AVOID"]
TIER = {"STRONG BUY": 5, "BUY": 4, "WATCH": 3, "NEUTRAL": 2, "AVOID": 1}
FILTER_OPTIONS = ["All Signals", "STRONG BUY + BUY"] + SIGNALS
DATA_COLS = ["Ticker", "Name", "CMP", "Chg1D", "Ret5D", "VolX", "POC", "VAH", "VAL", "AVWAP_Cat", "AVWAP_Earn",
             "AVWAP_Low", "SqueezeSpread", "Squeeze_Status", "Structural_Signal", "Score", "Signal", "LastBar", "Tier"]
DATA_CAPACITY = 800
MIN_SUCCESS_RATIO = 0.60          # refuse to overwrite the dashboard if fewer tickers than this loaded


# ----------------------------------------------------------------------------- constituents
def _read_csv_url(fname: str) -> pd.DataFrame | None:
    import requests
    for base in CSV_BASES:
        try:
            r = requests.get(base + fname, headers=UA, timeout=25)
            if r.status_code == 200 and b"Symbol" in r.content[:400]:
                from io import BytesIO
                return pd.read_csv(BytesIO(r.content))
        except Exception as e:  # network hiccup -> try next mirror
            LOG.debug("%s%s failed: %s", base, fname, e)
    return None


def load_constituents(only: list[str] | None = None):
    """Returns ({index: [(symbol, name), ...]}, [warnings])."""
    CACHE.mkdir(exist_ok=True)
    cache_f = CACHE / "constituents.json"
    cached = json.loads(cache_f.read_text()) if cache_f.exists() else {}
    out, warns = {}, []
    ov_f = HERE / "constituents_override.json"           # optional: {"Nifty Pharma": ["SUNPHARMA", "CIPLA"]}
    override = json.loads(ov_f.read_text()) if ov_f.exists() else {}
    for idx, fname in INDEX_FILES.items():
        if only and idx not in only:
            continue
        if override.get(idx):
            out[idx] = [(x.replace(".NS", "").upper(), x.replace(".NS", "").upper()) for x in override[idx]]
            warns.append(f"{idx}: using constituents_override.json")
            continue
        df = _read_csv_url(fname)
        if df is not None and "Symbol" in df.columns and len(df):
            name_col = "Company Name" if "Company Name" in df.columns else "Symbol"
            out[idx] = [(str(s).strip(), str(n).strip()) for s, n in zip(df["Symbol"], df[name_col])]
            cached[idx] = out[idx]
        elif idx in cached:
            out[idx] = [tuple(x) for x in cached[idx]]
            warns.append(f"{idx}: live download failed, used cached constituents")
        elif idx == "Nifty 50":
            out[idx] = [(s, s) for s in NIFTY50_FALLBACK]
            warns.append("Nifty 50: live download failed, used built-in list")
        else:
            warns.append(f"{idx}: could not be loaded (no live file, no cache)")
    cache_f.write_text(json.dumps(cached))
    return out, warns


# ----------------------------------------------------------------------------- prices
def _clean(df: pd.DataFrame) -> pd.DataFrame | None:
    if df is None or df.empty:
        return None
    df = df.copy()
    df.index = pd.to_datetime(df.index).tz_localize(None).normalize()
    need = ["Open", "High", "Low", "Close", "Volume"]
    if any(c not in df.columns for c in need):
        return None
    df = df[need].dropna(subset=["Close", "High", "Low"])
    df = df[(df["Close"] > 0) & (df["High"] >= df["Low"])]
    # drop a dead/partial last bar (zero volume) so a mid-session run cannot poison the signals
    if len(df) and df["Volume"].iloc[-1] <= 0:
        df = df.iloc[:-1]
    return df if len(df) else None


def download_prices(symbols: list[str], batch=25, retries=4):
    import yfinance as yf
    prices, failed = {}, []
    for i in range(0, len(symbols), batch):
        chunk = symbols[i:i + batch]
        raw = None
        for attempt in range(1, retries + 1):
            try:
                raw = yf.download(chunk, period="2y", interval="1d", auto_adjust=False,
                                  group_by="ticker", threads=6, progress=False)
                if raw is not None and not raw.empty:
                    break
            except Exception as e:
                LOG.warning("batch %d attempt %d failed: %s", i // batch, attempt, e)
            time.sleep(3 * attempt)                   # back-off grows on every retry
        for s in chunk:
            df = None
            if raw is not None and not raw.empty:
                try:
                    df = raw[s] if isinstance(raw.columns, pd.MultiIndex) else raw
                except KeyError:
                    df = None
            df = _clean(df)
            if df is None:
                failed.append(s)
            else:
                prices[s] = df
        LOG.info("downloaded %d/%d", min(i + batch, len(symbols)), len(symbols))
        time.sleep(1.2)                                  # be polite: avoids Yahoo rate limits
    # second pass: one-by-one for anything that failed
    still = []
    for s in failed:
        time.sleep(0.4)
        try:
            df = _clean(yf.Ticker(s).history(period="2y", interval="1d", auto_adjust=False))
        except Exception:
            df = None
        if df is None:
            still.append(s)
        else:
            prices[s] = df
    return prices, still


def synthetic_prices(symbols: list[str], seed=7):
    rng = np.random.default_rng(seed)
    idx = pd.bdate_range(end=pd.Timestamp.today().normalize(), periods=500)
    out = {}
    for s in symbols:
        p0 = rng.uniform(100, 5000)
        r = rng.normal(0.0004, 0.017, len(idx))
        c = p0 * np.exp(np.cumsum(r))
        o = c * (1 + rng.normal(0, 0.004, len(idx)))
        h = np.maximum(o, c) * (1 + np.abs(rng.normal(0, 0.007, len(idx))))
        l = np.minimum(o, c) * (1 - np.abs(rng.normal(0, 0.007, len(idx))))
        v = rng.uniform(1e5, 5e6, len(idx))
        out[s] = pd.DataFrame({"Open": o, "High": h, "Low": l, "Close": c, "Volume": v}, index=idx)
    return out


# ----------------------------------------------------------------------------- analytics
def volume_profile(win: pd.DataFrame, bins=BINS, va=VALUE_AREA):
    """Daily bars -> volume spread uniformly across each bar's high-low range."""
    lo, hi, vol = win["Low"].to_numpy(float), win["High"].to_numpy(float), win["Volume"].to_numpy(float)
    pmin, pmax = lo.min(), hi.max()
    if not pmax > pmin:
        return None
    edges = np.linspace(pmin, pmax, bins + 1)
    prof = np.zeros(bins)
    for l, h, v in zip(lo, hi, vol):
        if v <= 0:
            continue
        if h <= l:
            prof[min(max(np.searchsorted(edges, l, side="right") - 1, 0), bins - 1)] += v
            continue
        ov = np.clip(np.minimum(edges[1:], h) - np.maximum(edges[:-1], l), 0, None)
        s = ov.sum()
        if s > 0:
            prof += v * ov / s
    total = prof.sum()
    if total <= 0:
        return None
    poc_i = int(prof.argmax())
    lo_i = hi_i = poc_i
    acc = prof[poc_i]
    while acc < va * total and (lo_i > 0 or hi_i < bins - 1):
        up = prof[hi_i + 1] if hi_i < bins - 1 else -1.0
        dn = prof[lo_i - 1] if lo_i > 0 else -1.0
        if up >= dn:
            hi_i += 1; acc += prof[hi_i]
        else:
            lo_i -= 1; acc += prof[lo_i]
    return (edges[poc_i] + edges[poc_i + 1]) / 2, edges[hi_i + 1], edges[lo_i]


def avwap(df: pd.DataFrame, anchor) -> float:
    if anchor is None or pd.isna(anchor):
        return np.nan
    sub = df.loc[df.index >= pd.Timestamp(anchor)]
    if sub.empty or sub["Volume"].sum() <= 0:
        return np.nan
    tp = (sub["High"] + sub["Low"] + sub["Close"]) / 3
    return float((tp * sub["Volume"]).sum() / sub["Volume"].sum())


def fin(*xs) -> bool:
    return all(x is not None and np.isfinite(x) for x in xs)


def classify_structure(cmp_, prev, open_, low, poc, vah, val, spread):
    if fin(vah, val):
        if cmp_ > vah and prev <= vah:
            return "🚀 Value Area Breakout"
        if low <= val * 1.005 and cmp_ > val and cmp_ > open_:
            return "🛡️ Institutional Support Bounce"
        if (prev < val or prev > vah) and val <= cmp_ <= vah:
            return "↩️ Value Area Mean Reversion"
    if fin(spread) and spread < SQUEEZE_BAND:
        return "🧲 AVWAP Squeeze"
    if fin(vah, val):
        if cmp_ > vah:
            return "📈 Above Value Area - bullish acceptance"
        if cmp_ < val:
            return "📉 Below Value Area - distribution"
    return "⚖️ Inside Value Area - balanced"


def score_and_signal(r: dict):
    """0-100 score from whichever components are available (missing ones are excluded, not zeroed)."""
    cmp_, vah, val, poc = r["CMP"], r["VAH"], r["VAL"], r["POC"]
    comps = []                                           # (earned, possible)
    if fin(cmp_, vah, val) and vah > val:                # 1. position vs value area   (25)
        pos = (cmp_ - val) / (vah - val)
        comps.append((25 if pos > 1 else 17 if pos >= .5 else 9 if pos >= 0 else 0, 25))
    anchors = [a for a in (r["AVWAP_Cat"], r["AVWAP_Earn"], r["AVWAP_Low"]) if fin(a)]
    frac_above = np.nan
    if fin(cmp_) and anchors:                            # 2. above anchored VWAPs     (30)
        frac_above = sum(cmp_ > a for a in anchors) / len(anchors)
        comps.append((30 * frac_above, 30))
    v = r["VolX"]
    if fin(v):                                           # 3. volume vs 20d average    (15)
        comps.append((15 if v >= 2 else 12 if v >= 1.5 else 9 if v >= 1.2 else 6 if v >= 1 else 3 if v >= .7 else 0, 15))
    if fin(cmp_, poc):                                   # 4. above POC                (10)
        comps.append((10 if cmp_ > poc else 0, 10))
    sp = r["SqueezeSpread"]
    if fin(sp):                                          # 5. squeeze, direction-aware (10)
        comps.append(((10 if frac_above == 1 else 3) if sp < SQUEEZE_BAND else 0, 10))
    if fin(r["Ret5D"]):                                  # 6. 5-day momentum           (10)
        g = r["Ret5D"]
        comps.append((10 if g > .03 else 7 if g > .01 else 4 if g > 0 else 0, 10))
    if len(comps) < 3:
        return np.nan, "NEUTRAL"
    sc = 100 * sum(e for e, _ in comps) / sum(p for _, p in comps)
    if fin(cmp_, vah) and cmp_ > vah * 1.08:             # don't chase over-extended moves
        sc -= 10
    sc = round(float(np.clip(sc, 0, 100)), 2)
    vol_ok = (not fin(v)) or v >= 1.2
    if sc >= 75 and (frac_above == 1 or not fin(frac_above)) and vol_ok and (not fin(poc, cmp_) or cmp_ >= poc):
        sig = "STRONG BUY"
    elif sc >= 60 and (not fin(frac_above) or frac_above >= .66):
        sig = "BUY"
    elif sc >= 45:
        sig = "WATCH"
    elif sc >= 30:
        sig = "NEUTRAL"
    else:
        sig = "AVOID"
    return sc, sig


def squeeze_text(sp):
    return "" if not fin(sp) else f"{'Tight' if sp < SQUEEZE_BAND else 'Wide'} {sp * 100:.2f}%"


def analyse(sym: str, df: pd.DataFrame, p: dict):
    lb = int(p["lookback"])
    if len(df) < max(lb, 30):
        return None
    win = df.tail(lb)
    vp = volume_profile(win)
    if vp is None:
        return None
    poc, vah, val = vp
    close = df["Close"]
    cmp_, prev = float(close.iloc[-1]), float(close.iloc[-2])
    avg_v = df["Volume"].iloc[-(VOL_AVG_DAYS + 1):-1].mean()
    anch = [avwap(df, p["catalyst"]), avwap(df, p["earnings"]), avwap(df, p["swing_low"])]
    av = [a for a in anch if fin(a)]
    spread = (max(av) - min(av)) / np.mean(av) if len(av) >= 2 else np.nan
    row = dict(
        Ticker=sym, CMP=round(cmp_, 2), Chg1D=cmp_ / prev - 1,
        Ret5D=cmp_ / float(close.iloc[-6]) - 1 if len(close) > 6 else np.nan,
        VolX=round(float(df["Volume"].iloc[-1] / avg_v), 2) if avg_v and avg_v > 0 else np.nan,
        POC=round(poc, 2), VAH=round(vah, 2), VAL=round(val, 2),
        AVWAP_Cat=round(anch[0], 2) if fin(anch[0]) else np.nan,
        AVWAP_Earn=round(anch[1], 2) if fin(anch[1]) else np.nan,
        AVWAP_Low=round(anch[2], 2) if fin(anch[2]) else np.nan,
        SqueezeSpread=spread, LastBar=df.index[-1].date(),
    )
    row["Squeeze_Status"] = squeeze_text(spread)
    row["Structural_Signal"] = classify_structure(cmp_, prev, float(df["Open"].iloc[-1]),
                                                  float(df["Low"].iloc[-1]), poc, vah, val, spread)
    row["Score"], row["Signal"] = score_and_signal(row)
    return row


# ----------------------------------------------------------------------------- inputs / seed
def read_prev_inputs(path: Path) -> dict:
    p = dict(DEFAULTS)
    if not path.exists():
        return p
    try:
        from openpyxl import load_workbook
        ws = load_workbook(path)["Scanner"]
        def d(v):
            return v.date() if isinstance(v, dt.datetime) else v if isinstance(v, dt.date) else None
        mapping = {"index": "B5", "sort": "D5", "signal_filter": "F5", "lookback": "H5",
                   "catalyst": "J5", "earnings": "L5", "swing_low": "N5"}
        for k, ref in mapping.items():
            v = ws[ref].value
            if v in (None, ""):
                continue
            p[k] = int(v) if k == "lookback" else d(v) if k in ("catalyst", "earnings", "swing_low") else str(v)
            if p[k] is None:
                p[k] = DEFAULTS[k]
        if p["sort"] not in SORT_OPTIONS: p["sort"] = DEFAULTS["sort"]
        if p["signal_filter"] not in FILTER_OPTIONS: p["signal_filter"] = DEFAULTS["signal_filter"]
        if p["index"] != ALL and p["index"] not in INDEX_FILES: p["index"] = DEFAULTS["index"]
    except Exception as e:
        LOG.warning("could not read previous inputs (%s); using defaults", e)
    return p


def load_seed(path: str):
    """Import a snapshot produced by the original Quadratic scanner (Sheet1 layout)."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).worksheets[0]
    p = dict(DEFAULTS)
    lab = {"Volume Profile Lookback (days)": "lookback", "Catalyst Date (YYYY-MM-DD)": "catalyst",
           "Earnings Date (YYYY-MM-DD)": "earnings", "Swing Low Date (YYYY-MM-DD)": "swing_low"}
    hdr, rows = None, []
    for r in ws.iter_rows(values_only=True):
        if r[0] in lab and r[1] is not None:
            v = r[1]
            p[lab[r[0]]] = int(v) if lab[r[0]] == "lookback" else (v.date() if isinstance(v, dt.datetime) else v)
        if r[0] == "Ticker":
            hdr = list(r); continue
        if hdr and r[0] and str(r[0]).endswith(".NS"):
            rows.append(dict(zip(hdr, r)))
    out = []
    for x in rows:
        txt = str(x.get("Squeeze_Status") or "")
        try:
            sp = float(txt.split()[-1].strip("%")) / 100
        except Exception:
            sp = np.nan
        sig = str(x.get("Structural_Signal") or "")
        emoji = {"Value Area Breakout": "🚀 ", "Institutional Support Bounce": "🛡️ ", "Value Area Mean Reversion": "↩️ ",
                 "AVWAP Squeeze": "🧲 ", "Above Value Area": "📈 ", "Below Value Area": "📉 ", "Inside Value Area": "⚖️ "}
        for k, e in emoji.items():
            if sig.startswith(k):
                sig = e + sig
        row = dict(Ticker=x["Ticker"], CMP=x["CMP"], Chg1D=np.nan, Ret5D=np.nan, VolX=x["Volume_vs_Avg"],
                   POC=x["POC"], VAH=x["VAH"], VAL=x["VAL"], AVWAP_Cat=x["AVWAP_Catalyst"], AVWAP_Earn=np.nan,
                   AVWAP_Low=x["AVWAP_Low"], SqueezeSpread=sp, Squeeze_Status=txt, Structural_Signal=sig,
                   LastBar=None)
        row["Score"], row["Signal"] = score_and_signal(row)
        out.append(row)
    return out, p


# ----------------------------------------------------------------------------- publish (feeds Google Sheets)
def publish_csv(master: pd.DataFrame, membership: dict, meta: dict, params: dict, out_dir: str):
    """Writes data.csv (no header; same column order as the Data sheet) and meta.csv (9 rows, 1 column).
    The Google Sheet pulls both with IMPORTDATA. Written atomically."""
    d = Path(out_dir); d.mkdir(parents=True, exist_ok=True)
    index_names = list(INDEX_FILES.keys())
    rows = []
    for rec in master.to_dict("records")[:DATA_CAPACITY]:
        row = []
        for c in DATA_COLS:
            v = rec.get(c)
            if v is None or (isinstance(v, float) and not np.isfinite(v)):
                v = ""
            elif c == "LastBar":
                v = str(v)
            row.append(v)
        row += [1 if rec["Ticker"] in membership.get(nm, set()) else "" for nm in index_names]
        rows.append(row)
    ptxt = (f"lookback={params['lookback']}d, catalyst={params['catalyst']}, earnings={params['earnings']}, "
            f"swing-low={params['swing_low']}")
    metarows = [meta["run"], meta["source"], meta["asof"], meta.get("requested", ""), meta.get("loaded", ""),
                meta.get("failed_n", ""), ", ".join(meta.get("failed", []))[:3000] or "none",
                " | ".join(meta.get("warnings", []))[:3000] or "none", ptxt]
    for name, payload in (("data.csv", rows), ("meta.csv", [[x] for x in metarows])):
        tmp = d / (name + ".tmp")
        with open(tmp, "w", newline="", encoding="utf-8") as fh:
            csv.writer(fh).writerows(payload)
        os.replace(tmp, d / name)
    LOG.info("published %d rows to %s", len(rows), d)


# ----------------------------------------------------------------------------- workbook
def build_workbook(master: pd.DataFrame, membership: dict, params: dict, meta: dict, path: Path, import_mode: bool = False):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule, DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation

    FONT = "Arial"
    NAVY, INK = "1F3864", "1B2631"
    fill = lambda h: PatternFill("solid", start_color=h, end_color=h)
    f = lambda **k: Font(name=FONT, **k)
    thin = Side(style="thin", color="D5D8DC")
    med = Side(style="medium", color="1B2631")
    hair = Border(bottom=thin)

    index_names = list(INDEX_FILES.keys())
    n = len(master)
    CAP = max(800, n + 100)          # pre-built capacity so Apps Script / later runs can fill any index
    DASH = min(CAP, 600)             # dashboard rows
    last = CAP + 1
    wb = Workbook()
    ws = wb.active; ws.title = "Scanner"
    wb_b = wb.create_sheet("Index Breadth")
    wd = wb.create_sheet("Data")
    wg = wb.create_sheet("Guide")
    wr = wb.create_sheet("Run Log")
    wl = wb.create_sheet("Lists")
    for sh, col in ((ws, "1F3864"), (wb_b, "0E7C7B"), (wd, "7F8C8D"), (wg, "B9770E"), (wr, "7F8C8D"), (wl, "7F8C8D")):
        sh.sheet_properties.tabColor = col

    # ---------------- Lists
    wl["A1"], wl["C1"], wl["E1"] = "Index", "Sort By", "Signal Filter"
    for i, v in enumerate([ALL] + index_names, 2): wl.cell(i, 1, v)
    for i, v in enumerate(SORT_OPTIONS, 2): wl.cell(i, 3, v)
    for i, v in enumerate(FILTER_OPTIONS, 2): wl.cell(i, 5, v)
    for c in "ACE": wl[f"{c}1"].font = f(bold=True); wl.column_dimensions[c].width = 28
    wl["A1"].font = wl["C1"].font = wl["E1"].font = f(bold=True)
    wl.sheet_state = "hidden"
    idx_rng = f"=Lists!$A$2:$A${len(index_names) + 2}"
    sort_rng = f"=Lists!$C$2:$C${len(SORT_OPTIONS) + 1}"
    flt_rng = f"=Lists!$E$2:$E${len(FILTER_OPTIONS) + 1}"

    # ---------------- Data sheet
    cols = DATA_COLS
    C = {c: L(i + 1) for i, c in enumerate(cols)}                  # column letters
    m0 = len(cols) + 1                                               # first membership column index
    mcol = {nm: L(m0 + i) for i, nm in enumerate(index_names)}
    h0 = m0 + len(index_names)
    helpers = ["RowID", "InIdx", "SigOK", "Sel", "Key", "Rank"]
    H = {c: L(h0 + i) for i, c in enumerate(helpers)}
    m_first, m_last = mcol[index_names[0]], mcol[index_names[-1]]
    rng = lambda col: f"Data!${col}$2:${col}${last}"

    for j, c in enumerate(cols + index_names + helpers, 1):
        cell = wd.cell(1, j, c)
        cell.font = f(bold=True, color="FFFFFF")
        cell.fill = fill("34495E" if j < m0 else "117864" if j < h0 else "7F8C8D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wd.row_dimensions[1].height = 34
    recs_all = master.to_dict("records")
    for i in range(2, last + 1):
        rec = recs_all[i - 2] if i - 2 < len(recs_all) else None
        if rec:
            for j, c in enumerate(cols, 1):
                v = rec.get(c)
                if v is None or (isinstance(v, float) and not np.isfinite(v)):
                    v = None
                wd.cell(i, j, v).font = f(size=10)
            for nm in index_names:
                if rec["Ticker"] in membership.get(nm, set()):
                    wd[f"{mcol[nm]}{i}"] = 1
        r = i
        sort_choice = (f"CHOOSE(MATCH(Scanner!$D$5,Lists!$C$2:$C${len(SORT_OPTIONS) + 1},0),"
                       f"IF({C['Score']}{r}=\"\",-1E9,{C['Tier']}{r}*1000+{C['Score']}{r}),IF({C['VolX']}{r}=\"\",-1E9,{C['VolX']}{r}),"
                       f"IF({C['Chg1D']}{r}=\"\",-1E9,{C['Chg1D']}{r}),IF({C['Ret5D']}{r}=\"\",-1E9,{C['Ret5D']}{r}),"
                       f"IF({C['SqueezeSpread']}{r}=\"\",-1E9,-{C['SqueezeSpread']}{r}))")
        wd[f"{H['RowID']}{r}"] = r - 1
        wd[f"{H['InIdx']}{r}"] = (f'=IF(Scanner!$B$5="{ALL}",IF($A{r}<>"",1,0),IFERROR(INDEX(${m_first}{r}:${m_last}{r},'
                                  f'MATCH(Scanner!$B$5,${m_first}$1:${m_last}$1,0))+0,0))')
        wd[f"{H['SigOK']}{r}"] = (f'=IF(Scanner!$F$5="All Signals",1,IF(Scanner!$F$5="STRONG BUY + BUY",'
                                  f'IF(OR({C["Signal"]}{r}="STRONG BUY",{C["Signal"]}{r}="BUY"),1,0),'
                                  f'IF({C["Signal"]}{r}=Scanner!$F$5,1,0)))')
        wd[f"{H['Sel']}{r}"] = f"={H['InIdx']}{r}*{H['SigOK']}{r}"
        wd[f"{H['Key']}{r}"] = f'=IF({H["Sel"]}{r}=1,ROUND({sort_choice},6),"")'
        wd[f"{H['Rank']}{r}"] = (f'=IF({H["Sel"]}{r}=1,COUNTIFS({rng(H["Sel"])},1,{rng(H["Key"])},">"&{H["Key"]}{r})'
                                 f'+COUNTIFS({rng(H["Sel"])},1,{rng(H["Key"])},{H["Key"]}{r},{rng(H["RowID"])},"<"&{H["RowID"]}{r})+1,"")')
        for c in helpers:
            wd[f"{H[c]}{r}"].font = f(size=9, color="7F8C8D")
        for c, fmt in (("CMP", "#,##0.00"), ("Chg1D", "0.00%"), ("Ret5D", "0.00%"), ("VolX", "0.00"), ("POC", "#,##0.00"),
                       ("VAH", "#,##0.00"), ("VAL", "#,##0.00"), ("AVWAP_Cat", "#,##0.00"), ("AVWAP_Earn", "#,##0.00"),
                       ("AVWAP_Low", "#,##0.00"), ("SqueezeSpread", "0.00%"), ("Score", "0.0"), ("Tier", "0"), ("LastBar", "yyyy-mm-dd")):
            wd[f"{C[c]}{r}"].number_format = fmt
    if import_mode:      # Google Sheets pulls the published CSV into A2 and spills it down/right
        wd["A2"] = ("=IFERROR(IMPORTDATA('Run Log'!$B$15&\"data.csv?v=\"&TEXT(NOW(),\"yyyymmddhh\")),\"\")")
    wd.freeze_panes = "C2"
    wd.auto_filter.ref = f"A1:{H['Rank']}{last}"
    for j in range(1, h0 + len(helpers)):
        wd.column_dimensions[L(j)].width = 13
    wd.column_dimensions["B"].width = 30; wd.column_dimensions[C["Structural_Signal"]].width = 36
    for nm in index_names:
        wd.column_dimensions[mcol[nm]].width = 11

    # ---------------- Scanner dashboard
    widths = {"A": 6, "B": 15, "C": 26, "D": 15, "E": 8, "F": 11, "G": 9, "H": 9, "I": 9, "J": 11, "K": 11, "L": 11,
              "M": 12, "N": 12, "O": 12, "P": 14, "Q": 38, "R": 2, "S": 6}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.column_dimensions["S"].hidden = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    for rr, hgt in {1: 34, 2: 17, 3: 5, 4: 15, 5: 25, 6: 5, 7: 15, 8: 30, 9: 5, 10: 17, 11: 32}.items():
        ws.row_dimensions[rr].height = hgt

    ws.merge_cells("A1:Q1")
    ws["A1"] = "⚡ NIFTY STRUCTURAL SCANNER  ·  Volume Profile + Multi-Anchor AVWAP + Signal Ranking"
    ws["A1"].font = f(size=18, bold=True, color="FFFFFF"); ws["A1"].fill = fill(NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.merge_cells("A2:Q2")
    ws["A2"] = ("='Run Log'!B5&\"  ·  Source: \"&'Run Log'!B4&\"  ·  \"&IF(SUM(" + rng(H["InIdx"]) + ")=0,"
                "\"⚠ No stocks loaded for this index in the current run - pick another index or re-run the scanner\","
                "SUM(" + rng(H["InIdx"]) + ")&\" stocks in selected index  ·  \"&SUM(" + rng(H["Sel"]) + ")&\" shown after signal filter\")")
    ws["A2"].font = f(size=9, italic=True, color="FFFFFF"); ws["A2"].fill = fill("2E4A7D")
    ws["A2"].alignment = Alignment(vertical="center", indent=1)

    # control tiles (label row 4, value row 5)
    tiles = [("B", "C", "INDEX SELECTION ▼", params["index"], None, idx_rng),
             ("D", "E", "SORT BY ▼", params["sort"], None, sort_rng),
             ("F", "G", "SIGNAL FILTER ▼", params["signal_filter"], None, flt_rng),
             ("H", "I", "LOOKBACK (DAYS)", int(params["lookback"]), "0", None),
             ("J", "K", "CATALYST DATE", params["catalyst"], "yyyy-mm-dd", None),
             ("L", "M", "EARNINGS DATE", params["earnings"], "yyyy-mm-dd", None),
             ("N", "O", "SWING-LOW DATE", params["swing_low"], "yyyy-mm-dd", None)]
    for a, b, lab, val, fmt, dv_rng in tiles:
        ws.merge_cells(f"{a}4:{b}4"); ws.merge_cells(f"{a}5:{b}5")
        ws[f"{a}4"] = lab; ws[f"{a}4"].font = f(size=8, bold=True, color="FFFFFF")
        ws[f"{a}4"].fill = fill("1F3864" if dv_rng else "7D6608"); ws[f"{a}4"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"{a}5"] = val; ws[f"{a}5"].font = f(size=11, bold=True, color="0000FF")
        ws[f"{a}5"].alignment = Alignment(horizontal="center", vertical="center")
        for col in (a, b):
            ws[f"{col}5"].fill = fill("FFF6DA")
            ws[f"{col}5"].border = Border(left=med if col == a else None, right=med if col == b else None, top=med, bottom=med)
        if fmt: ws[f"{a}5"].number_format = fmt
        if dv_rng:
            dv = DataValidation(type="list", formula1=dv_rng, allow_blank=False, showErrorMessage=True,
                                errorTitle="Invalid choice", error="Pick a value from the dropdown list.")
            ws.add_data_validation(dv); dv.add(f"{a}5")
    dv_lb = DataValidation(type="whole", operator="between", formula1="20", formula2="250", showErrorMessage=True,
                           error="Lookback must be a whole number of days between 20 and 250.")
    dv_dt = DataValidation(type="date", operator="greaterThan", formula1="36526", showErrorMessage=True, error="Enter a valid date.")
    ws.add_data_validation(dv_lb); ws.add_data_validation(dv_dt)
    dv_lb.add("H5"); [dv_dt.add(c) for c in ("J5", "L5", "N5")]

    # KPI tiles (row 7 label / row 8 value)
    sigr, idr = rng(C["Signal"]), rng(H["InIdx"])
    kpis = [("B", "C", "IN VIEW", f"=SUM({rng(H['Sel'])})", "1F3864", "0"),
            ("D", "E", "STRONG BUY", f'=COUNTIFS({idr},1,{sigr},"STRONG BUY")', "117A44", "0"),
            ("F", "G", "BUY", f'=COUNTIFS({idr},1,{sigr},"BUY")', "27AE60", "0"),
            ("H", "I", "WATCH", f'=COUNTIFS({idr},1,{sigr},"WATCH")', "D4A017", "0"),
            ("J", "K", "NEUTRAL", f'=COUNTIFS({idr},1,{sigr},"NEUTRAL")', "7F8C8D", "0"),
            ("L", "M", "AVOID", f'=COUNTIFS({idr},1,{sigr},"AVOID")', "C0392B", "0"),
            ("N", "O", "AVG SCORE", f'=IFERROR(AVERAGEIFS({rng(C["Score"])},{idr},1),"-")', "6C3483", "0.0"),
            ("P", "Q", "LAST RUN (IST)", "='Run Log'!B3", "34495E", "General")]
    for a, b, lab, val, colr, fmt in kpis:
        ws.merge_cells(f"{a}7:{b}7"); ws.merge_cells(f"{a}8:{b}8")
        ws[f"{a}7"] = lab; ws[f"{a}7"].font = f(size=8, bold=True, color="FFFFFF")
        ws[f"{a}8"] = val; ws[f"{a}8"].number_format = fmt
        ws[f"{a}8"].font = f(size=18 if lab != "LAST RUN (IST)" else 12, bold=True, color=colr)
        for rr in (7, 8):
            ws[f"{a}{rr}"].alignment = Alignment(horizontal="center", vertical="center")
        for col in (a, b):
            ws[f"{col}7"].fill = fill(colr); ws[f"{col}8"].fill = fill("F4F6F7")
            ws[f"{col}8"].border = Border(bottom=Side(style="thick", color=colr))

    # column group headers (row 10) and column headers (row 11)
    groups = [("A", "B", "C", "IDENTITY", "1F3864"), ("D", "D", "E", "SIGNAL", "117A44"),
              ("F", "F", "I", "PRICE & MOMENTUM", "2E5C9A"), ("J", "J", "L", "VOLUME PROFILE", "B9770E"),
              ("M", "M", "P", "ANCHORED VWAPs", "6C3483"), ("Q", "Q", "Q", "STRUCTURE", "0E6655")]
    heads = {"A": "RANK", "B": "TICKER", "C": "COMPANY", "D": "SIGNAL", "E": "SCORE", "F": "CMP", "G": "1D %", "H": "5D %",
             "I": "VOL ×AVG", "J": "POC", "K": "VAH", "L": "VAL", "M": "AVWAP\nCATALYST", "N": "AVWAP\nEARNINGS",
             "O": "AVWAP\nSWING LOW", "P": "SQUEEZE", "Q": "STRUCTURAL SIGNAL"}
    first_row, last_row = 12, 11 + DASH
    group_start = {g[0] for g in groups}
    for first, _, lastc, label, colr in groups:
        ws.merge_cells(f"{first}10:{lastc}10")
        ws[f"{first}10"] = label; ws[f"{first}10"].font = f(size=9, bold=True, color="FFFFFF")
        ws[f"{first}10"].alignment = Alignment(horizontal="center", vertical="center")
        for ci in range(ord(first), ord(lastc) + 1):
            col = chr(ci)
            ws[f"{col}10"].fill = fill(colr)
            ws[f"{col}11"].fill = fill(colr); ws[f"{col}11"].value = heads[col]
            ws[f"{col}11"].font = f(size=9, bold=True, color="FFFFFF")
            ws[f"{col}11"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if col in group_start:
                ws[f"{col}10"].border = ws[f"{col}11"].border = Border(left=Side(style="medium", color="FFFFFF"))

    src = {"B": "Ticker", "C": "Name", "D": "Signal", "E": "Score", "F": "CMP", "G": "Chg1D", "H": "Ret5D", "I": "VolX",
           "J": "POC", "K": "VAH", "L": "VAL", "M": "AVWAP_Cat", "N": "AVWAP_Earn", "O": "AVWAP_Low", "P": "Squeeze_Status",
           "Q": "Structural_Signal"}
    fmts = {"E": "0.0", "F": "#,##0.00;-#,##0.00;;@", "G": "+0.00%;-0.00%;0.00%", "H": "+0.00%;-0.00%;0.00%", "I": '0.00"x"',
            "J": "#,##0.00;-#,##0.00;;@", "K": "#,##0.00;-#,##0.00;;@", "L": "#,##0.00;-#,##0.00;;@",
            "M": "#,##0.00;-#,##0.00;;@", "N": "#,##0.00;-#,##0.00;;@", "O": "#,##0.00;-#,##0.00;;@"}
    aligns = {"A": "center", "B": "left", "C": "left", "D": "center", "E": "center", "P": "center", "Q": "left"}
    for r in range(first_row, last_row + 1):
        ws.row_dimensions[r].height = 19
        ws[f"S{r}"] = f'=IFERROR(MATCH(ROW()-{first_row - 1},{rng(H["Rank"])},0),"")'
        ws[f"A{r}"] = f'=IF($S{r}="","",ROW()-{first_row - 1})'
        for col, key in src.items():
            base = f"INDEX({rng(C[key])},$S{r})"
            ws[f"{col}{r}"] = (f'=IF($S{r}="","",IF({base}="","",{base}))' if col in "GHNPQ" or key in ("Name",)
                               else f'=IF($S{r}="","",{base})')
        for col in "ABCDEFGHIJKLMNOPQ":
            c = ws[f"{col}{r}"]
            c.font = f(size=10, bold=col in "BDE")
            c.alignment = Alignment(horizontal=aligns.get(col, "right"), vertical="center", indent=1 if col in "CQ" else 0)
            if col in fmts: c.number_format = fmts[col]
            c.border = Border(bottom=thin, left=med if col in group_start and col != "A" else None)
        ws[f"Q{r}"].border = Border(bottom=thin, right=med, left=med)
    ws.freeze_panes = "C12"
    ws.auto_filter.ref = None

    T = f"{first_row}:{last_row}"
    rg = lambda col: f"{col}{first_row}:{col}{last_row}"
    # --- signal colours
    for text, bg, fg in (("STRONG BUY", "117A44", "FFFFFF"), ("BUY", "ABEBC6", "145A32"), ("WATCH", "FCF3CF", "7D6608"),
                         ("NEUTRAL", "E5E8E8", "424949"), ("AVOID", "F5B7B1", "922B21")):
        ws.conditional_formatting.add(rg("D"), CellIsRule(operator="equal", formula=[f'"{text}"'], fill=fill(bg),
                                                          font=Font(name=FONT, bold=True, color=fg)))
    ws.conditional_formatting.add(rg("E"), ColorScaleRule(start_type="num", start_value=0, start_color="F1948A",
                                                          mid_type="num", mid_value=50, mid_color="F9E79F",
                                                          end_type="num", end_value=100, end_color="58D68D"))
    for col in "GH":
        ws.conditional_formatting.add(rg(col), CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, color="117A44", bold=True)))
        ws.conditional_formatting.add(rg(col), CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C0392B", bold=True)))
    ws.conditional_formatting.add(rg("I"), FormulaRule(formula=[f"AND(ISNUMBER(I{first_row}),I{first_row}>=1.5)"], fill=fill("F7DC6F"), font=Font(name=FONT, bold=True)))
    ws.conditional_formatting.add(rg("I"), FormulaRule(formula=[f"AND(ISNUMBER(I{first_row}),I{first_row}>=1.2)"], fill=fill("FCF3CF")))
    # CMP vs value area
    ws.conditional_formatting.add(rg("F"), FormulaRule(formula=[f"AND(ISNUMBER(F{first_row}),F{first_row}>K{first_row})"], fill=fill("ABEBC6"), font=Font(name=FONT, bold=True, color="145A32")))
    ws.conditional_formatting.add(rg("F"), FormulaRule(formula=[f"AND(ISNUMBER(F{first_row}),F{first_row}<L{first_row})"], fill=fill("F5B7B1"), font=Font(name=FONT, bold=True, color="922B21")))
    # AVWAPs: green = price above anchor (support), red = price below (resistance)
    for col in "MNO":
        ws.conditional_formatting.add(rg(col), FormulaRule(formula=[f"AND(ISNUMBER({col}{first_row}),{col}{first_row}>0,$F{first_row}>{col}{first_row})"], fill=fill("D5F5E3")))
        ws.conditional_formatting.add(rg(col), FormulaRule(formula=[f"AND(ISNUMBER({col}{first_row}),{col}{first_row}>0,$F{first_row}<={col}{first_row})"], fill=fill("FADBD8")))
    ws.conditional_formatting.add(rg("J"), FormulaRule(formula=[f"AND(ISNUMBER(J{first_row}),$F{first_row}>J{first_row})"], fill=fill("EAF2F8")))
    ws.conditional_formatting.add(rg("P"), FormulaRule(formula=[f'LEFT(P{first_row},5)="Tight"'], fill=fill("D7BDE2"), font=Font(name=FONT, bold=True, color="4A235A")))
    for kw, bg, fg in (("Breakout", "ABEBC6", "145A32"), ("Support Bounce", "ABEBC6", "145A32"), ("Mean Reversion", "AED6F1", "1B4F72"),
                       ("Squeeze", "D7BDE2", "4A235A"), ("bullish", "D5F5E3", "1E8449"), ("distribution", "FADBD8", "922B21"),
                       ("balanced", "EAEDED", "566573")):
        ws.conditional_formatting.add(rg("Q"), FormulaRule(formula=[f'ISNUMBER(SEARCH("{kw}",Q{first_row}))'], fill=fill(bg), font=Font(name=FONT, color=fg)))
    # zebra banding last (lowest priority)
    ws.conditional_formatting.add(f"A{first_row}:Q{last_row}", FormulaRule(formula=[f'AND($S{first_row}<>"",MOD(ROW(),2)=0)'], fill=fill("F2F6FA")))
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "10:11"

    # ---------------- Index Breadth
    wb_b.sheet_view.showGridLines = False
    wb_b.merge_cells("A1:J1")
    wb_b["A1"] = "🧭 INDEX BREADTH  ·  which indices are loaded with the most BUY / STRONG BUY setups"
    wb_b["A1"].font = f(size=16, bold=True, color="FFFFFF"); wb_b["A1"].fill = fill("0E7C7B")
    wb_b["A1"].alignment = Alignment(vertical="center", indent=1); wb_b.row_dimensions[1].height = 32
    wb_b.merge_cells("A2:J2")
    wb_b["A2"] = "All figures are live formulas over the Data sheet. % Bullish = (STRONG BUY + BUY) / stocks loaded."
    wb_b["A2"].font = f(size=9, italic=True, color="566573")
    bh = ["INDEX", "STOCKS", "STRONG BUY", "BUY", "WATCH", "NEUTRAL", "AVOID", "% BULLISH", "AVG SCORE", "RANK"]
    bcol = ["0E7C7B", "0E7C7B", "117A44", "27AE60", "D4A017", "7F8C8D", "C0392B", "6C3483", "6C3483", "34495E"]
    for j, (t, colr) in enumerate(zip(bh, bcol), 1):
        c = wb_b.cell(4, j, t); c.font = f(bold=True, color="FFFFFF", size=10); c.fill = fill(colr)
        c.alignment = Alignment(horizontal="center", vertical="center")
    wb_b.row_dimensions[4].height = 24
    brow = {}
    for i, nm in enumerate([ALL] + index_names, 5):
        brow[nm] = i
        if nm == ALL:
            cond = None
            wb_b.cell(i, 2, f"=COUNTA({rng(C['Ticker'])})")
        else:
            cond = rng(mcol[nm])
            wb_b.cell(i, 2, f"=SUM({cond})")
        wb_b.cell(i, 1, nm)
        for j, s in zip(range(3, 8), SIGNALS):
            wb_b.cell(i, j, f'=COUNTIFS({sigr},"{s}")' if cond is None else f'=COUNTIFS({cond},1,{sigr},"{s}")')
        wb_b.cell(i, 8, f'=IF(B{i}=0,"",(C{i}+D{i})/B{i})')
        wb_b.cell(i, 9, f'=IF(B{i}=0,"",AVERAGE({rng(C["Score"])}))' if cond is None else
                  f'=IFERROR(AVERAGEIFS({rng(C["Score"])},{cond},1),"")')
        bl = 4 + len(index_names) + 1
        wb_b.cell(i, 10, f'=IF(H{i}="","",RANK(H{i},$H$6:$H${bl}))' if cond else "")
        for j in range(1, 11):
            c = wb_b.cell(i, j); c.font = f(size=10, bold=(j == 1 or nm == ALL)); c.border = hair
            c.alignment = Alignment(horizontal="left" if j == 1 else "center", indent=1 if j == 1 else 0)
            if nm == ALL: c.fill = fill("EAF2F8")
        wb_b.cell(i, 8).number_format = "0.0%"; wb_b.cell(i, 9).number_format = "0.0"
    bl = 4 + len(index_names) + 1
    wb_b.conditional_formatting.add(f"H5:H{bl}", ColorScaleRule(start_type="num", start_value=0, start_color="F1948A", mid_type="num", mid_value=0.25, mid_color="F9E79F", end_type="num", end_value=0.6, end_color="58D68D"))
    wb_b.conditional_formatting.add(f"I5:I{bl}", ColorScaleRule(start_type="num", start_value=20, start_color="F1948A", mid_type="num", mid_value=50, mid_color="F9E79F", end_type="num", end_value=80, end_color="58D68D"))
    wb_b.conditional_formatting.add(f"C5:C{bl}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill("ABEBC6"), font=Font(name=FONT, bold=True, color="145A32")))
    wb_b.conditional_formatting.add(f"D5:D{bl}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill("D5F5E3")))
    wb_b.conditional_formatting.add(f"G5:G{bl}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill("FADBD8")))
    for j, wdt in zip("ABCDEFGHIJ", (30, 9, 12, 8, 9, 10, 9, 11, 11, 7)):
        wb_b.column_dimensions[j].width = wdt
    ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "% Bullish (STRONG BUY + BUY) by index"
    ch.add_data(Reference(wb_b, min_col=8, min_row=4, max_row=bl), titles_from_data=True)
    ch.set_categories(Reference(wb_b, min_col=1, min_row=5, max_row=bl))
    ch.legend = None; ch.height, ch.width = 17, 16; ch.y_axis.number_format = "0%"
    ch.y_axis.delete = False; ch.x_axis.delete = False; ch.x_axis.scaling.orientation = "maxMin"
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
    ch.series[0].graphicalProperties.solidFill = "0E7C7B"
    wb_b.add_chart(ch, "L4")
    wb_b.freeze_panes = "A5"

    # ---------------- Run log
    wr.sheet_view.showGridLines = False
    wr["A1"] = "RUN LOG"; wr["A1"].font = f(size=14, bold=True, color=NAVY)
    rows = [("Run time (IST)", meta["run"]), ("Data source", meta["source"]), ("Latest price bar in data", meta["asof"]),
            ("Tickers requested", meta.get("requested")), ("Tickers loaded", meta.get("loaded")),
            ("Tickers failed", meta.get("failed_n")), ("Failed tickers", ", ".join(meta.get("failed", [])) or "none"),
            ("Constituent warnings", " | ".join(meta.get("warnings", [])) or "none"),
            ("Parameters used", f"lookback={params['lookback']}d, catalyst={params['catalyst']}, earnings={params['earnings']}, swing-low={params['swing_low']}")]
    for i, (k, v) in enumerate(rows, 3):
        wr.cell(i, 1, k).font = f(bold=True); wr.cell(i, 2, None if import_mode else v).font = f()
        wr.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top"); wr.cell(i, 1).alignment = Alignment(vertical="top")
    if import_mode:
        wr["B3"] = ("=IFERROR(IMPORTDATA('Run Log'!$B$15&\"meta.csv?v=\"&TEXT(NOW(),\"yyyymmddhh\")),"
                    "\"Not loaded - paste your GitHub data URL in B15\")")
        wr["A14"] = "SETUP"; wr["A14"].font = f(bold=True, color=NAVY)
        wr["A15"] = "GitHub data URL  (yellow)"; wr["A15"].font = f(bold=True)
        wr["B15"] = "https://raw.githubusercontent.com/YOUR-USER/YOUR-REPO/main/data/"
        wr["B15"].font = f(color="0000FF", bold=True); wr["B15"].fill = fill("FFF6DA")
        wr["A16"] = "How it works"; wr["A16"].font = f(bold=True)
        wr["B16"] = ("GitHub Actions runs the scanner every weekday evening and publishes data.csv + meta.csv. "
                     "The Data sheet (cell A2) and this sheet (B3) read them with IMPORTDATA. "
                     "Keep the URL ending in /data/. The repository must be public.")
        wr["B16"].alignment = Alignment(wrap_text=True, vertical="top")
    wr.column_dimensions["A"].width = 30; wr.column_dimensions["B"].width = 120

    # ---------------- Guide
    wg.sheet_view.showGridLines = False
    wg.column_dimensions["A"].width = 26; wg.column_dimensions["B"].width = 110
    wg["A1"] = "GUIDE · HOW THE SCANNER THINKS"; wg["A1"].font = f(size=14, bold=True, color=NAVY)
    guide = [
        ("HOW TO USE", ""),
        ("Index Selection", "Pick any Nifty index. Rows re-rank instantly (formula-driven; no macro needed)."),
        ("Sort By / Signal Filter", "Default ranks STRONG BUY above BUY above WATCH..., then by Score. Or re-sort by Volume, 1D %, 5D % or Squeeze tightness; filter to STRONG BUY, BUY, etc."),
        ("Yellow cells", "Inputs. Lookback / Catalyst / Earnings / Swing-Low dates are read by the scanner on its NEXT run."),
        ("Cell colours", "CMP green = above VAH, red = below VAL. AVWAP cells green = price above that anchor, red = below."),
        ("", ""),
        ("SCORE (0-100)", "Components that are missing are excluded and the rest re-scaled, never silently counted as zero."),
        ("Value-area position  25", "Above VAH 25 · upper half of value area 17 · lower half 9 · below VAL 0"),
        ("Anchored VWAPs  30", "10 points per anchor (Catalyst, Earnings, Swing-Low) that price is trading above"),
        ("Volume vs 20d avg  15", ">=2x 15 · >=1.5x 12 · >=1.2x 9 · >=1x 6 · >=0.7x 3 · else 0"),
        ("Above POC  10", "10 if CMP is above the point of control"),
        ("Squeeze  10", "AVWAPs within 2%: 10 if price is above all of them (bullish coil), 3 otherwise, 0 if wide"),
        ("5-day momentum  10", ">3% 10 · >1% 7 · >0% 4 · else 0"),
        ("Over-extension penalty", "-10 if CMP is more than 8% above VAH (avoid chasing)"),
        ("", ""),
        ("SIGNALS", ""),
        ("STRONG BUY", "Score >= 75, price above ALL anchors and POC, volume >= 1.2x average"),
        ("BUY", "Score >= 60 and price above at least 2 of 3 anchors"),
        ("WATCH / NEUTRAL / AVOID", "Score 45-60 / 30-45 / below 30"),
        ("", ""),
        ("STRUCTURE LABELS", "Priority order: Breakout > Support Bounce > Mean Reversion > AVWAP Squeeze > Above / Below / Inside value area"),
        ("Breakout", "Close above VAH after previous close at or below VAH"),
        ("Support Bounce", "Day's low within 0.5% of VAL, closes above VAL and above open"),
        ("Mean Reversion", "Previous close outside the value area, today closes back inside it"),
        ("", ""),
        ("ASSUMPTIONS", ""),
        ("Volume profile", "Built from daily bars: each day's volume is spread evenly across its high-low range over 50 price bins; value area = 70% of volume around the POC."),
        ("Volume vs Avg", "Latest volume / mean of the previous 20 sessions."),
        ("Squeeze spread", "(max - min) / mean of the available AVWAPs; Tight if below 2%."),
        ("Refresh", "Scheduled after market close (see nifty_scanner.py / scanner_eod.yml). Prices: Yahoo Finance via yfinance; constituents: NSE Indices."),
        ("Disclaimer", "Rule-based technical screen on public, delayed data. Not investment advice. Verify before trading."),
    ]
    for i, (a, b) in enumerate(guide, 3):
        wg.cell(i, 1, a).font = f(bold=True, color=NAVY if b == "" and a else "000000", size=11 if b == "" and a else 10)
        wg.cell(i, 2, b).font = f(size=10); wg.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
        wg.cell(i, 1).alignment = Alignment(vertical="top")
        if b == "" and a:
            wg.cell(i, 1).fill = wg.cell(i, 2).fill = fill("EAF2F8")

    wb.active = 0
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)       # atomic: a crashed run can never leave a half-written dashboard
    LOG.info("wrote %s", path)


# ----------------------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=str(HERE / "NIFTY_STRUCTURAL_SCANNER.xlsx"))
    ap.add_argument("--demo", action="store_true", help="synthetic prices, no network (self-test)")
    ap.add_argument("--seed-xlsx", help="import a snapshot from the old Quadratic-style Sheet1 layout")
    ap.add_argument("--publish-dir", help="also write data.csv / meta.csv here (for the Google Sheet to import)")
    ap.add_argument("--template-live", action="store_true", help="build the empty Google-Sheets workbook that imports the CSVs")
    ap.add_argument("--only", help="comma-separated index names to restrict the run")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if a.verbose else logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    out = Path(a.out)
    ist = dt.datetime.now(dt.timezone(dt.timedelta(hours=5, minutes=30)))
    meta = dict(run=ist.strftime("%Y-%m-%d %H:%M"), warnings=[])

    if a.template_live:
        params = read_prev_inputs(out)
        empty = pd.DataFrame(columns=DATA_COLS)
        meta.update(source="", asof="", requested="", loaded="", failed_n="", failed=[])
        build_workbook(empty, {}, params, meta, out, import_mode=True)
        return 0

    if a.seed_xlsx:
        recs, params = load_seed(a.seed_xlsx)
        prev = read_prev_inputs(out)  # keep any dropdown choices made in the new file
        master = pd.DataFrame(recs)
        master["Name"] = master["Ticker"].str.replace(".NS", "", regex=False)
        membership = {"Nifty 50": set(master["Ticker"])}
        meta.update(source="Seed snapshot imported from the original Quadratic scanner file (not a live pull)",
                    asof="Snapshot", requested=len(master), loaded=len(master), failed_n=0, failed=[],
                    warnings=["Only Nifty 50 is populated in this seed. Run the live scanner to fill every other index."])
    else:
        params = read_prev_inputs(out)
        only = [s.strip() for s in a.only.split(",")] if a.only else None
        cons, warns = load_constituents(only) if not a.demo else (
            {"Nifty 50": [(s, s) for s in NIFTY50_FALLBACK],
             "Nifty Bank": [(s, s) for s in "HDFCBANK ICICIBANK KOTAKBANK AXISBANK SBIN INDUSINDBK".split()],
             "Nifty IT": [(s, s) for s in "TCS INFY WIPRO HCLTECH TECHM".split()]}, [])
        names, membership = {}, {}
        for idx, lst in cons.items():
            membership[idx] = {s + ".NS" for s, _ in lst}
            for s, nm in lst: names[s + ".NS"] = nm
        universe = sorted(names)
        LOG.info("universe: %d tickers across %d indices", len(universe), len(cons))
        prices, failed = (synthetic_prices(universe), []) if a.demo else download_prices(universe)
        if len(prices) < MIN_SUCCESS_RATIO * max(len(universe), 1):
            LOG.error("only %d/%d tickers loaded - keeping the previous dashboard untouched", len(prices), len(universe))
            return 2
        recs = [r for s, df in prices.items() if (r := analyse(s, df, params))]
        master = pd.DataFrame(recs)
        master["Name"] = master["Ticker"].map(names)
        latest = max(df.index[-1] for df in prices.values()).date()
        meta.update(source="Yahoo Finance (yfinance), NSE Indices constituents" if not a.demo else "SYNTHETIC DEMO DATA",
                    asof=f"{latest:%Y-%m-%d}", requested=len(universe), loaded=len(recs),
                    failed_n=len(universe) - len(recs), failed=sorted(set(universe) - set(master["Ticker"])), warnings=warns)
        meta["asof"] = f"Prices as of {meta['asof']}"
    master["Tier"] = master["Signal"].map(TIER)
    master = master.sort_values("Ticker").reset_index(drop=True)
    for k in ("Chg1D", "Ret5D", "AVWAP_Earn"):
        if k not in master: master[k] = np.nan
    if a.publish_dir:
        publish_csv(master, membership, meta, params, a.publish_dir)
    build_workbook(master, membership, params, meta, out)
    sig = master["Signal"].value_counts().to_dict()
    LOG.info("signals: %s", {s: sig.get(s, 0) for s in SIGNALS})
    return 0


if __name__ == "__main__":
    sys.exit(main())
